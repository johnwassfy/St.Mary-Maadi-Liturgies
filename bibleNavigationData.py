"""Book/chapter/verse navigation data for bibleWindow.

Builds a per-book JSON cache mapping chapter -> {start_slide, verses: {verse: slide}}
by scanning each Phase 3 deck's normalized chapter titles and superscripted verse runs.
"""
import json
import os
import re
from pathlib import Path

from pptx import Presentation

from commonFunctions import relative_path
from superscripting import BOOK_SHORTCUTS, BOOK_SECTION_NAMES

OLD_TESTAMENT = "عهد قديم"
NEW_TESTAMENT = "عهد جديد"
TESTAMENTS = (OLD_TESTAMENT, NEW_TESTAMENT)

# Runs with this baseline offset are verse-number superscripts (see superscripting.py).
VERSE_BASELINE = "40000"
ARABIC_TO_ASCII_DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩", "0123456789")

# A handful of verse numbers (e.g. Psalm 119:101+) were left as plain embedded digits by
# the original superscripting pass, which only handled up to 2-digit numbers. This pattern
# recognizes a digit run butted directly against the following word (no space), matching
# how the superscripted verse-number runs are always immediately followed by verse text.
_EMBEDDED_VERSE_PATTERN = re.compile(r"(?<![0-9\u0660-\u0669])([0-9\u0660-\u0669]+)(?=\S)")

# Bump whenever _build_manifest's detection logic changes, so stale caches (keyed only on
# source mtime) still get rebuilt.
_MANIFEST_VERSION = 2

_CACHE_DIR = Path(relative_path(r"Data\CopyData\bible_navigation"))


class BibleNavigationError(Exception):
    """Raised when a book's chapter/verse manifest cannot be built or loaded."""


def _phase3_folder(testament):
    return Path(relative_path(os.path.join("الكتاب المقدس", testament, "phase3")))


def _book_files(testament):
    """Map book number to its Phase 3 file path for one testament."""
    files = {}
    folder = _phase3_folder(testament)
    if not folder.is_dir():
        return files
    for path in folder.glob("*.pptx"):
        match = re.match(r"\s*(\d+)", path.stem)
        if match:
            files[int(match.group(1))] = path
    return files


def list_books():
    """Return ordered book entries: (testament, number, arabic_name, shortcut, path_or_None)."""
    books = []
    for testament in TESTAMENTS:
        files = _book_files(testament)
        names = BOOK_SECTION_NAMES[testament]
        shortcuts = BOOK_SHORTCUTS[testament]
        for number in sorted(names):
            books.append((
                testament,
                number,
                names[number],
                shortcuts.get(number),
                files.get(number),
            ))
    return books


def _cache_path(testament, number):
    safe_testament = "OT" if testament == OLD_TESTAMENT else "NT"
    return _CACHE_DIR / f"{safe_testament}_{number:02d}.json"


def _verse_number_from_run(text):
    digits = text.translate(ARABIC_TO_ASCII_DIGITS).strip()
    return int(digits) if digits.isdigit() else None


def _extract_chapter_start(slide):
    """Return the chapter number if this slide is a chapter's first slide."""
    for shape in slide.shapes:
        if shape.is_placeholder and shape.placeholder_format.idx == 0 and shape.has_text_frame:
            text = shape.text.strip()
            match = re.search(r"(\d+)\s*$", text) if text else None
            return int(match.group(1)) if match else None
    return None


def _extract_verse_numbers(slide, expected_verse):
    """Return (verse_numbers_found, next_expected_verse) for a slide.

    Superscripted runs are trusted unconditionally. Non-superscripted runs are only
    checked against the expected next verse number, to safely catch the rare unformatted
    verse markers (see _EMBEDDED_VERSE_PATTERN) without misreading ordinary numbers.
    """
    found = []
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                r_pr = run.font._rPr
                if r_pr is not None and r_pr.get("baseline") == VERSE_BASELINE:
                    verse_number = _verse_number_from_run(run.text)
                    if verse_number is not None:
                        found.append(verse_number)
                        expected_verse = verse_number + 1
                    continue
                for match in _EMBEDDED_VERSE_PATTERN.finditer(run.text):
                    verse_number = _verse_number_from_run(match.group(1))
                    if verse_number == expected_verse:
                        found.append(verse_number)
                        expected_verse += 1
    return found, expected_verse


def _build_manifest(path):
    presentation = Presentation(path)
    chapters = {}
    current_chapter = None
    expected_verse = 1
    for slide_number, slide in enumerate(presentation.slides, start=1):
        if slide_number == 1:
            continue  # index/navigation slide, no chapter content
        chapter_number = _extract_chapter_start(slide)
        if chapter_number is not None:
            current_chapter = chapter_number
            chapters.setdefault(
                str(current_chapter), {"start_slide": slide_number, "verses": {}}
            )
            expected_verse = 1
        if current_chapter is None:
            continue
        verse_numbers, expected_verse = _extract_verse_numbers(slide, expected_verse)
        verses = chapters[str(current_chapter)]["verses"]
        for verse_number in verse_numbers:
            verses.setdefault(str(verse_number), slide_number)

    if not chapters:
        raise BibleNavigationError(f"No chapter-start slides found in {path.name}")
    return {
        "version": _MANIFEST_VERSION,
        "source_mtime": path.stat().st_mtime,
        "chapters": chapters,
    }


def get_book_manifest(testament, number, path):
    """Load a cached manifest for one book, rebuilding it if the source file changed.

    Returns {"chapters": {"1": {"start_slide": int, "verses": {"1": int, ...}}, ...}}.
    Chapter and verse keys are strings (JSON round-trip); callers should str() lookups.
    """
    cache_file = _cache_path(testament, number)
    try:
        source_mtime = path.stat().st_mtime
    except OSError as error:
        raise BibleNavigationError(f"Cannot read {path}: {error}") from error

    if cache_file.is_file():
        try:
            cached = json.loads(cache_file.read_text(encoding="utf-8"))
            if (
                cached.get("source_mtime") == source_mtime
                and cached.get("version") == _MANIFEST_VERSION
            ):
                return cached
        except (json.JSONDecodeError, OSError):
            pass

    manifest = _build_manifest(path)
    _CACHE_DIR.mkdir(parents=True, exist_ok=True)
    cache_file.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
    return manifest
